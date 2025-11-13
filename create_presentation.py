#!/usr/bin/env python3
"""
PowerPoint Presentation Generator for Almarai International Business Report
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from docx import Document

def extract_document_content(docx_path):
    """Extract content from DOCX file"""
    doc = Document(docx_path)
    content = []
    for para in doc.paragraphs:
        if para.text.strip():
            content.append(para.text.strip())
    return content

def add_title_slide(prs, title, subtitle):
    """Add title slide"""
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_shape = slide.shapes.title
    subtitle_shape = slide.placeholders[1]
    
    title_shape.text = title
    subtitle_shape.text = subtitle
    
    # Style title
    title_frame = title_shape.text_frame
    title_frame.paragraphs[0].font.size = Pt(44)
    title_frame.paragraphs[0].font.bold = True
    title_frame.paragraphs[0].font.color.rgb = RGBColor(0, 51, 102)

def add_content_slide(prs, title, content_items, layout_type=1):
    """Add content slide with bullet points"""
    slide = prs.slides.add_slide(prs.slide_layouts[layout_type])
    
    # Set title
    title_shape = slide.shapes.title
    title_shape.text = title
    title_shape.text_frame.paragraphs[0].font.size = Pt(32)
    title_shape.text_frame.paragraphs[0].font.bold = True
    title_shape.text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 51, 102)
    
    # Add content
    if len(slide.placeholders) > 1:
        content_shape = slide.placeholders[1]
        text_frame = content_shape.text_frame
        text_frame.clear()
        
        for item in content_items:
            p = text_frame.add_paragraph()
            p.text = item
            p.level = 0
            p.font.size = Pt(16)
            p.space_after = Pt(12)

def add_table_slide(prs, title, headers, rows):
    """Add slide with table"""
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    
    # Set title
    title_shape = slide.shapes.title
    title_shape.text = title
    title_shape.text_frame.paragraphs[0].font.size = Pt(32)
    title_shape.text_frame.paragraphs[0].font.bold = True
    title_shape.text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 51, 102)
    
    # Add table
    rows_count = len(rows) + 1
    cols_count = len(headers)
    left = Inches(0.5)
    top = Inches(2)
    width = Inches(9)
    height = Inches(0.5)
    
    table = slide.shapes.add_table(rows_count, cols_count, left, top, width, height).table
    
    # Set headers
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = header
        cell.text_frame.paragraphs[0].font.bold = True
        cell.text_frame.paragraphs[0].font.size = Pt(14)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0, 51, 102)
        cell.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    
    # Set rows
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.text = str(cell_text)
            cell.text_frame.paragraphs[0].font.size = Pt(12)

def create_presentation():
    """Create the complete PowerPoint presentation"""
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    # Slide 1: Title Slide
    add_title_slide(
        prs,
        "Almarai's Expansion into Indonesia",
        "International Market Entry Report\nMBA 6112 International Business"
    )
    
    # Slide 2: Executive Summary
    add_content_slide(
        prs,
        "Executive Summary",
        [
            "Almarai Company: Saudi Arabia's leading dairy and food producer",
            "Objective: Identify optimal foreign market and entry strategy",
            "Analysis Framework: PESTLE analysis, comparative factor weighting, entry strategy evaluation",
            "Selected Market: Indonesia (270M+ population, Halal-certified demand)",
            "Recommended Entry: Joint venture with local food/logistics partner",
            "Strategy: Localized marketing, pricing, and CSR for sustainable integration"
        ]
    )
    
    # Slide 3: Company Background
    add_content_slide(
        prs,
        "Company Background",
        [
            "Founded: 1977 in Riyadh, Saudi Arabia",
            "Position: Largest vertically integrated dairy & food company in Middle East",
            "Operations: Dairy, juice, bakery, poultry, infant nutrition",
            "Scale: 200,000+ dairy cattle, 20 manufacturing facilities",
            "Revenue: SAR 18 billion (USD 4.8 billion) annually",
            "Vision: Consumer's preferred choice through quality, innovation, nutrition",
            "Alignment: Saudi Arabia's Vision 2030 for economic diversification"
        ]
    )
    
    # Slide 4: Drivers of International Expansion
    add_content_slide(
        prs,
        "Drivers of International Expansion",
        [
            "Market Saturation: Mature GCC markets require new growth avenues",
            "Vision 2030 Alignment: Export diversification and foreign investment",
            "Economic Diversification: Reduce dependence on regional dairy sales",
            "Cultural Compatibility: Indonesia's Muslim-majority aligns with Halal brand",
            "Regional Integration: ASEAN & RCEP membership reduces tariffs",
            "Core Competency: Leverage cold-chain logistics expertise",
            "Brand Globalization: Position as trusted global Halal brand"
        ]
    )
    
    # Slide 5: Alternative Markets Considered
    add_content_slide(
        prs,
        "Alternative Foreign Markets",
        [
            "India:",
            "  • Large scale but dominated by domestic cooperatives (Amul, Mother Dairy)",
            "  • High competition, strict regulations, limited cold-chain infrastructure",
            "  • Lower attractiveness for entry",
            "",
            "Indonesia:",
            "  • 80%+ dairy import dependency",
            "  • Strong economic growth (GDP ~USD 1.6 trillion)",
            "  • Young, urbanizing population",
            "  • Cultural and religious similarity with Saudi Arabia",
            "  • Selected for detailed assessment"
        ]
    )
    
    # Slide 6: PESTLE Analysis - Political & Economic
    add_content_slide(
        prs,
        "PESTLE Analysis: Political & Economic",
        [
            "Political:",
            "  • Stable democracy with consistent economic policy",
            "  • Government promotes food security and Halal industry",
            "  • Regulatory complexity in import licensing and certification",
            "",
            "Economic:",
            "  • GDP: USD 1.6 trillion; growth 5-5.5% annually",
            "  • Middle-class expansion: 140M by 2030",
            "  • Dairy market: 80% imports, 8-10% annual growth",
            "  • Currency risk: Rupiah volatility requires hedging"
        ]
    )
    
    # Slide 7: PESTLE Analysis - Social & Technological
    add_content_slide(
        prs,
        "PESTLE Analysis: Social & Technological",
        [
            "Social/Cultural:",
            "  • Largest Muslim population globally - Halal certification mandatory",
            "  • Preference for sweetened/flavored milk and yogurt drinks",
            "  • Family-oriented culture values brand trust and quality",
            "  • 95%+ literacy; highest social media usage globally",
            "",
            "Technological:",
            "  • E-commerce penetration >70%",
            "  • Growing logistics tech startups",
            "  • Underdeveloped cold-chain outside major cities",
            "  • Opportunity for Almarai's logistics expertise"
        ]
    )
    
    # Slide 8: PESTLE Analysis - Legal & Environmental
    add_content_slide(
        prs,
        "PESTLE Analysis: Legal & Environmental",
        [
            "Legal:",
            "  • Comprehensive consumer and Halal laws (BPOM, MUI)",
            "  • High but transparent compliance costs",
            "  • Improving IP protection under WTO-TRIPS",
            "",
            "Environmental:",
            "  • Deforestation and carbon-emission concerns",
            "  • Sustainability pressure on dairy waste and packaging",
            "  • Government incentives for green manufacturing",
            "  • Renewable energy opportunities"
        ]
    )
    
    # Slide 9: Market Comparison Table
    add_table_slide(
        prs,
        "India vs Indonesia Comparison",
        ["Criteria", "India", "Indonesia", "Weight"],
        [
            ["Political Stability", "Moderate", "High", "0.1"],
            ["Market Growth", "High", "High", "0.2"],
            ["Cultural Fit", "Low", "High", "0.15"],
            ["Competition", "Very High", "Moderate", "0.15"],
            ["Infrastructure", "Moderate", "Developing", "0.1"],
            ["Regulatory Ease", "Complex", "Moderate", "0.15"],
            ["Import Dependency", "Low", "High (80%)", "0.15"]
        ]
    )
    
    # Slide 10: Final Country Selection
    add_content_slide(
        prs,
        "Why Indonesia?",
        [
            "High sustainable demand: Urbanization, rising income, Halal norms",
            "Strategic location: ASEAN & RCEP access to 2B consumers",
            "Political & economic stability: Market-based system",
            "Cultural compatibility: Saudi values enhance brand trust",
            "Favorable demographics: Young, health-conscious, tech-savvy",
            "Aligns with Saudi Vision 2030 objectives",
            "Opportunity for sustainable international growth"
        ]
    )
    
    # Slide 11: Entry Strategy Assessment
    add_content_slide(
        prs,
        "Entry Strategy Options",
        [
            "1. Exporting: Low investment but high transport costs, limited control",
            "2. Licensing/Franchising: Low risk but IP leakage risk, quality issues",
            "3. Joint Venture (RECOMMENDED):",
            "   • Shared risk with local partner (Indofood CBP, Mayora Indah)",
            "   • Access to distribution network and regulatory facilitation",
            "   • Local market insights with maintained quality control",
            "4. Wholly Owned Subsidiary: Full control but high investment & risk"
        ]
    )
    
    # Slide 12: Implementation Plan
    add_content_slide(
        prs,
        "Implementation Plan",
        [
            "Phase 1: Partner Selection & Due Diligence (6 months)",
            "  • Identify and evaluate potential JV partners",
            "  • Negotiate terms and establish governance structure",
            "",
            "Phase 2: Market Entry & Setup (12 months)",
            "  • Obtain Halal certification and regulatory approvals",
            "  • Establish distribution network and cold-chain logistics",
            "",
            "Phase 3: Launch & Scale (18-24 months)",
            "  • Product launch in major cities (Jakarta, Surabaya, Bandung)",
            "  • Digital marketing campaigns and brand awareness",
            "  • Expand to secondary markets based on performance"
        ]
    )
    
    # Slide 13: Marketing Strategy
    add_content_slide(
        prs,
        "Marketing & Brand Awareness Strategy",
        [
            "Brand Positioning: Premium Halal dairy from Middle East",
            "Target Segments: Urban middle-class families, health-conscious consumers",
            "Digital Marketing: Social media campaigns, influencer partnerships",
            "Localization: Adapt flavors to Indonesian preferences (sweetened, flavored)",
            "Retail Presence: Modern trade (supermarkets, convenience stores)",
            "E-commerce: Partner with Tokopedia, Shopee, GoFood",
            "Sampling & Promotions: In-store tastings, bundle offers"
        ]
    )
    
    # Slide 14: Pricing & CSR Strategy
    add_content_slide(
        prs,
        "Pricing & CSR Strategy",
        [
            "Pricing Strategy:",
            "  • Premium positioning (10-15% above local brands)",
            "  • Value packs for family consumption",
            "  • Promotional pricing during Ramadan and Eid",
            "",
            "CSR & Sustainability:",
            "  • Support local dairy farmers through training programs",
            "  • Eco-friendly packaging and waste reduction",
            "  • Community nutrition education initiatives",
            "  • Renewable energy in manufacturing facilities"
        ]
    )
    
    # Slide 15: Challenges & Solutions
    add_content_slide(
        prs,
        "Challenges & Solutions",
        [
            "Challenge: Complex regulatory environment",
            "Solution: Partner with local experts, invest in compliance team",
            "",
            "Challenge: Underdeveloped cold-chain infrastructure",
            "Solution: Leverage Almarai's logistics expertise, invest in facilities",
            "",
            "Challenge: Currency volatility (Rupiah)",
            "Solution: Financial hedging, local sourcing where possible",
            "",
            "Challenge: Competition from established local brands",
            "Solution: Differentiate through quality, Halal certification, premium positioning"
        ]
    )
    
    # Slide 16: Key Findings
    add_content_slide(
        prs,
        "Study Key Findings",
        [
            "Indonesia offers optimal market conditions for Almarai's expansion",
            "Strong cultural and religious alignment enhances brand acceptance",
            "Joint venture balances risk, control, and local market access",
            "Growing middle class drives demand for premium dairy products",
            "Digital infrastructure enables cost-effective marketing",
            "Regulatory challenges manageable with local partnership",
            "Sustainability and CSR critical for long-term success"
        ]
    )
    
    # Slide 17: Strategic Recommendations
    add_content_slide(
        prs,
        "Strategic Recommendations",
        [
            "1. Establish joint venture with reputable Indonesian food company",
            "2. Prioritize Halal certification and regulatory compliance",
            "3. Invest in cold-chain logistics infrastructure",
            "4. Launch with localized product portfolio (flavored milk, yogurt)",
            "5. Implement aggressive digital marketing strategy",
            "6. Build strong CSR program focused on sustainability",
            "7. Monitor currency risk and implement hedging strategies",
            "8. Plan phased expansion: major cities first, then secondary markets"
        ]
    )
    
    # Slide 18: Conclusion
    add_content_slide(
        prs,
        "Conclusion",
        [
            "Indonesia represents a strategic opportunity for Almarai",
            "Market conditions align with company's strengths and Vision 2030",
            "Joint venture approach minimizes risk while maximizing market access",
            "Success requires:",
            "  • Strong local partnership",
            "  • Cultural adaptation and localization",
            "  • Investment in infrastructure and compliance",
            "  • Commitment to sustainability and CSR",
            "Expected outcome: Sustainable growth and global brand recognition"
        ]
    )
    
    # Slide 19: Thank You
    add_title_slide(
        prs,
        "Thank You",
        "Questions & Discussion"
    )
    
    # Save presentation
    output_path = '/vercel/sandbox/Almarai_Indonesia_Expansion_Presentation.pptx'
    prs.save(output_path)
    print(f"✓ Presentation created successfully: {output_path}")
    return output_path

if __name__ == "__main__":
    create_presentation()
